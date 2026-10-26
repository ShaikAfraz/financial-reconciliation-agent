"""
Script to generate a Word Document (.docx) for the Non-Technical Guide.
"""
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tc_pr.append(tc_mar)

def create_non_tech_docx(output_path: str):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("The Autonomous Financial Reconciliation Agent: In Plain English")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(17)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(10)
    r_sub = p_sub.add_run("A Non-Technical Guide for Evaluators, Judges, Friends & General Audiences")
    r_sub.font.size = Pt(10.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # 30-sec pitch box
    tbl_pitch = doc.add_table(rows=1, cols=1)
    tbl_pitch.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pitch.autofit = False
    tbl_pitch.columns[0].width = Inches(7.0)
    c_p = tbl_pitch.rows[0].cells[0]
    set_cell_background(c_p, "F0F9FF")
    set_cell_margins(c_p, top=120, bottom=120, left=160, right=160)
    p_p = c_p.paragraphs[0]
    r_p_bold = p_p.add_run("The 30-Second Elevator Pitch: ")
    r_p_bold.font.bold = True
    r_p_bold.font.size = Pt(9.5)
    r_p_text = p_p.add_run(
        "Most budgeting apps act like dumb sticker machines: they guess a label for your purchase (like 'Groceries'), and then stop. If they make a mistake, they secretly ruin your budget numbers without you ever knowing. Our project is an intelligent digital accountant. Instead of just guessing, it cleans up messy bank descriptions, checks your spending limits, sounds the alarm on fraud or double-charges, and if you're over-budget, automatically figures out where to safely move money from (like cutting your entertainment budget to cover dinner) while strictly protecting essential bills like rent and healthcare."
    )
    r_p_text.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_h(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)
        return h

    # Section 1
    add_h("1. The Real-World Problem: Why Current Banking Apps Fail")
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(4)
    p1.add_run(
        "Whenever you swipe your card, the bank records a cryptic line of text like 'SQ *BLUE BOTTLE COFFEE #492 SFO CA'. Existing apps try to categorize this in a single guess. But money behaves like an interconnected web of envelopes. If you spend $120 on dinner and the app accidentally labels it 'Groceries': (1) Your Dining budget looks untouched, hiding your overspend; (2) Your Groceries envelope looks blown, raising false alarms; and (3) Automated advice tells you to stop buying groceries when the real problem was dining out."
    )

    # Section 2
    add_h("2. How Our Prototype Works: The 5 Steps of a Human Accountant")
    tbl_steps = doc.add_table(rows=6, cols=3)
    tbl_steps.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_steps.autofit = False
    col_w = [Inches(1.8), Inches(3.2), Inches(2.0)]
    for row in tbl_steps.rows:
        for idx, w in enumerate(col_w):
            row.cells[idx].width = w

    headers = ["Step & Tool", "What It Does (Simple Terms)", "Everyday Analogy"]
    for idx, h_text in enumerate(headers):
        c = tbl_steps.rows[0].cells[idx]
        set_cell_background(c, "1E293B")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    data_steps = [
        ("1. Categorize Transaction", "Cleans bank gibberish, extracts true merchant (Blue Bottle), assigns category with confidence score.", "Reading a wrinkled receipt and writing 'Coffee' on top."),
        ("2. Check Budget Rule", "Checks monthly limit for that envelope (Healthy, Warning at 80%, Exceeded over 100%).", "Counting how many dollar bills are left in your dining envelope."),
        ("3. Flag Anomaly", "Catches 3x price spikes, duplicate charges within minutes, or shady crypto merchants.", "Clerk saying: 'Wait, you were charged twice in 3 minutes!'"),
        ("4. Calculate Variance", "Examines all budget envelopes across the entire month to see total surpluses and deficits.", "Spreading all envelopes on the table to see the total scoreboard."),
        ("5. Propose Reallocation", "Moves spare cash from fun pools (Entertainment) to cover deficits—strictly protecting rent.", "Borrowing $20 from your movie jar to cover dinner, leaving rent untouched.")
    ]

    for r_idx, row_data in enumerate(data_steps, start=1):
        for c_idx, val in enumerate(row_data):
            c = tbl_steps.rows[r_idx].cells[c_idx]
            if r_idx % 2 == 1:
                set_cell_background(c, "F8FAFC")
            set_cell_margins(c, top=60, bottom=60, left=100, right=100)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Section 3
    add_h("3. Why Is This Smarter Than a Normal Chatbot (ChatGPT)?")
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_after = Pt(4)
    p3.add_run(
        "Chatbots like ChatGPT are notorious for hallucinating arithmetic—making up numbers or producing math that doesn't balance. Our prototype uses a Compound AI Architecture: (1) The AI is the Brain: it only handles reasoning and planning; (2) The Tools are the Hands: every database query and calculation is executed by hardcoded, un-bribeable mathematical code; and (3) The Executive Audit Trail shows every single step in plain English so you can verify it."
    )

    # Section 4
    add_h("4. Our Key Scientific Discovery: The 'Domino Effect' & Dollar Drift")
    p4 = doc.add_paragraph()
    p4.paragraph_format.space_after = Pt(4)
    p4.add_run(
        "When we deliberately injected an error into Step 1 across 11 scenarios, 100% of downstream tools blindly inherited the mistake, and 72.7% of final advice was compromised. Crucially, a misclassification creates double the financial distortion: mislabeling $100 distorts both the source and target envelope, causing $200 in Dollar Drift (Drift = 2X). Across our tests, simple errors created $6,353.08 in invisible budget distortion."
    )

    # Section 5
    add_h("5. The Ready-to-Use 60-Second Script")
    tbl_script = doc.add_table(rows=1, cols=1)
    tbl_script.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_script.autofit = False
    tbl_script.columns[0].width = Inches(7.0)
    c_s = tbl_script.rows[0].cells[0]
    set_cell_background(c_s, "F8FAFC")
    set_cell_margins(c_s, top=100, bottom=100, left=140, right=140)
    p_s = c_s.paragraphs[0]
    r_s = p_s.add_run(
        '"Have you ever used a budgeting app that wrongly tagged a dinner receipt as groceries and messed up all your charts? I built an autonomous AI financial reconciliation agent that fixes this problem. Instead of just slapping a single tag on a purchase, my agent acts like a live personal accountant: it parses the bank descriptor, checks your actual monthly budget limits, scans for fraud or double-billings, checks the overall financial health across all categories, and if you are in the red, it automatically moves spare money from fun categories like entertainment to balance your budget—without ever touching your rent or utility bills. Best of all, I proved mathematically that when regular AI apps make a small classification mistake, the error doubles across your ledger (Dollar Drift). My prototype uses a locked 5-tool pipeline to eliminate AI math hallucinations and keep budgets balanced."'
    )
    r_s.font.italic = True
    r_s.font.size = Pt(9.5)

    doc.save(output_path)
    print(f"Non-technical docx built at: {output_path}")

if __name__ == "__main__":
    out_docx_proj = "/Users/shaikmohammedafrazhussain/.gemini/antigravity-ide/scratch/financial-reconciliation-agent/non_technical_prototype_guide.docx"
    out_docx_art = "/Users/shaikmohammedafrazhussain/.gemini/antigravity-ide/brain/255a685a-4b23-4b66-b4d9-b2bbde53f572/non_technical_prototype_guide.docx"
    create_non_tech_docx(out_docx_proj)
    create_non_tech_docx(out_docx_art)
