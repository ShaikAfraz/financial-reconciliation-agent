"""
Generates a formal, professional Word Document (.docx) for the Future Scalability Plan & Architectural Roadmap.
"""
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tc_pr.append(tc_mar)

def create_scalability_docx(output_path: str):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("Future Scalability Plan & Architectural Roadmap")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Autonomous Financial Reconciliation Agent: Production & Capstone Final Submission Strategy")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(51, 65, 85)

    # Executive Summary Box
    tbl_exec = doc.add_table(rows=1, cols=1)
    tbl_exec.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_exec.autofit = False
    tbl_exec.columns[0].width = Inches(7.0)
    c_exec = tbl_exec.rows[0].cells[0]
    set_cell_background(c_exec, "F1F5F9")
    set_cell_margins(c_exec, top=120, bottom=120, left=160, right=160)
    p_exec = c_exec.paragraphs[0]
    r_exec_bold = p_exec.add_run("Executive Vision— ")
    r_exec_bold.font.bold = True
    r_exec_bold.font.size = Pt(9.5)
    r_exec_text = p_exec.add_run(
        "While the prototype proves the viability of an autonomous ReAct loop with a locked 5-tool catalog (achieving 100% Plan-Sequence Correctness and 94.9% Recommendation Groundedness), enterprise readiness requires scaling beyond localized SQLite batches and heuristic greedy rebalancing. This document outlines the comprehensive engineering and algorithmic blueprint for transforming the prototype into an enterprise-grade, publication-worthy final capstone submission. Key transitions include real-time streaming ingestion via Kafka, semantic vector memory with pgvector, formal Constrained Convex Optimization (Linear Programming via SciPy), self-correcting causal backtracking, and zero-trust PII masking."
    )
    r_exec_text.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)
        return h

    # Section 1
    add_h1("1. Architectural Evolution: From Prototype to Enterprise")
    p_arch = doc.add_paragraph()
    p_arch.paragraph_format.space_after = Pt(4)
    p_arch.add_run(
        "The production architecture migrates the standalone Python ReAct orchestrator into a decoupled, event-driven compound AI system. The pipeline separates high-throughput transaction ingestion, distributed relational/vector storage, multi-agent specialized worker execution, and immutable cryptographic compliance auditing."
    )

    # Section 2
    add_h1("2. The 5 Strategic Scalability Pillars")

    pillars = [
        ("Pillar I: Data Ingestion & Streaming Infrastructure",
         "Migrate from embedded SQLite batch loading to real-time ISO 20022 Open Banking webhooks (Plaid/MX/Yodlee). Decouple ingestion via an event-driven Kafka message bus (transactions.incoming) backed by a hybrid database tier: PostgreSQL for relational ACID accounts and TimescaleDB for sub-millisecond rolling category spend aggregations."),
        
        ("Pillar II: Semantic Vector Memory & Merchant Disambiguation",
         "Replace regex cleaning and token-set fuzzy matching with persistent semantic vector memory (pgvector). Embed 250,000+ verified national merchant descriptors via text-embedding-004. Calibrate confidence scores mathematically using cosine distance; any similarity below the ambiguity threshold (0.60) triggers an explicit clarification request rather than an ungrounded guess."),
        
        ("Pillar III: Mathematical Optimization Engine (Linear Programming)",
         "Upgrade from a heuristic greedy allocation to formal Constrained Convex Optimization (Linear Programming via SciPy/PuLP). The optimization minimizes surplus-draw penalties subject to strict budget-conservation constraints: sum(x_ij) <= Surplus_j, sum(x_ij) + u_i = Deficit_i, and hard floor protection (Housing, Healthcare, Utilities x_ij = 0). This eliminates sub-optimal transfers across multi-deficit scenarios."),
        
        ("Pillar IV: Self-Correcting Causal Guardrails & Backtracking",
         "Directly resolve the core failure mode discovered in the Error Propagation Analysis (where a Step 1 misclassification cascades into 100% of downstream tools and produces $6,353.08 in Dollar Drift). Introduce active post-step confidence gates: if Step 1 confidence is below 85%, execution halts and prompts the user. Additionally, a Conservation of Capital invariant (sum Delta Budgets == 0) ensures corrupt recommendations never reach the client."),
        
        ("Pillar V: Enterprise Privacy, Security & Cryptographic Auditing",
         "Implement a Zero-Trust PII redaction proxy (Microsoft Presidio) to mask PANs, SSNs, and customer names before prompts reach external LLM endpoints. Implement an immutable HMAC-SHA256 chained audit ledger ensuring non-repudiation for financial compliance (SOC2, GLBA), coupled with Human-in-the-Loop (HITL) approval gates for reallocations exceeding $250.00.")
    ]

    for p_title_text, p_body_text in pillars:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(1)
        bp.paragraph_format.space_after = Pt(3)
        r_pt = bp.add_run(f"{p_title_text}: ")
        r_pt.font.bold = True
        r_pt.font.size = Pt(9.5)
        r_pb = bp.add_run(p_body_text)
        r_pb.font.size = Pt(9.5)

    # Section 3
    add_h1("3. Phased Execution Roadmap for Final Submission")

    tbl_road = doc.add_table(rows=5, cols=4)
    tbl_road.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_road.autofit = False
    road_widths = [Inches(1.5), Inches(2.3), Inches(2.0), Inches(1.2)]
    for row in tbl_road.rows:
        for idx, width in enumerate(road_widths):
            row.cells[idx].width = width

    road_headers = ["Phase & Timeline", "Key Deliverables", "Target Tech Stack", "Estimated Effort"]
    for idx, h in enumerate(road_headers):
        c = tbl_road.rows[0].cells[idx]
        set_cell_background(c, "1E293B")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    road_data = [
        ("Phase 1: Semantic Memory (Weeks 1–2)", "• pgvector integration for 250k merchants\n• 10,000+ transaction dataset ingestion\n• Confidence calibration layer", "PostgreSQL, pgvector, sentence-transformers", "10–14 days"),
        ("Phase 2: Math Optimization (Weeks 3–4)", "• SciPy / PuLP Linear Programming solver\n• Multi-deficit convex rebalancing\n• Zero-hallucination constraint proofs", "SciPy, PuLP, Python 3.11, NumPy", "10–12 days"),
        ("Phase 3: Self-Correction & HITL (Weeks 5–6)", "• Step 1 confidence rollback gates\n• PII tokenization via Presidio\n• Push approval triggers (> $250)", "Microsoft Presidio, WebSockets, HMAC Ledger", "10–14 days"),
        ("Phase 4: Dashboard & Defense (Weeks 7–8)", "• Next.js 14 high-density executive ledger UI\n• Heuristic vs. LP comparative benchmark\n• Final thesis paper, slide deck & demo video", "Next.js 14, Tremor, Docker Compose, Chart.js", "12–14 days")
    ]

    for r_idx, data in enumerate(road_data, start=1):
        for c_idx, text in enumerate(data):
            cell = tbl_road.rows[r_idx].cells[c_idx]
            if r_idx % 2 == 1:
                set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Section 4
    add_h1("4. Architectural Comparison: Prototype vs. Capstone Final")

    tbl_comp = doc.add_table(rows=7, cols=3)
    tbl_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_comp.autofit = False
    comp_widths = [Inches(1.8), Inches(2.6), Inches(2.6)]
    for row in tbl_comp.rows:
        for idx, width in enumerate(comp_widths):
            row.cells[idx].width = width

    comp_headers = ["Technical Dimension", "Current Working Prototype", "Future Scalable Capstone Submission"]
    for idx, h in enumerate(comp_headers):
        c = tbl_comp.rows[0].cells[idx]
        set_cell_background(c, "1E293B")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    comp_data = [
        ("Data Layer", "Embedded SQLite (reconciliation.db)", "Distributed PostgreSQL + TimescaleDB + pgvector"),
        ("Merchant Categorization", "Regex rules + fuzzy string matching", "Semantic Vector RAG with calibrated uncertainty"),
        ("Reallocation Engine", "Heuristic greedy surplus rebalancing", "Constrained Linear Programming (SciPy / PuLP)"),
        ("Error Resilience", "Linear forward execution (cascading drift)", "Active confidence gating, backtracking & invariants"),
        ("Security & Governance", "Local database storage", "Presidio PII masking + HMAC-SHA256 chained audit"),
        ("Client & Human Controls", "Localhost Browser Console", "Real-time WebSockets, push alerts & HITL gates")
    ]

    for r_idx, data in enumerate(comp_data, start=1):
        for c_idx, text in enumerate(data):
            cell = tbl_comp.rows[r_idx].cells[c_idx]
            if r_idx % 2 == 1:
                set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True

    # Section 5
    add_h1("5. Capstone Defense Strategy (The 3-Act Demonstration)")
    p_def = doc.add_paragraph()
    p_def.paragraph_format.space_after = Pt(0)
    p_def.add_run(
        "For the final evaluation committee, present the submission using a compelling 3-Act structure: "
        "(1) Act 1: The Status Quo Fragility: Demonstrate how single-turn budgeting tools fail by corrupting dependent envelopes when categorizations err. "
        "(2) Act 2: The Working Agentic Prototype: Showcase the 5-tool ReAct loop live (Scenario 06) with explainable audit cards, SQL aggregations, and balanced rebalancing. "
        "(3) Act 3: The Empirical Frontier & Production Rigor: Present the Plan-Sequence Benchmark (100%), the Dollar Drift findings ($6,353.08), and demonstrate how the proposed Linear Programming solver and self-correction guardrails resolve the fundamental failure modes of compound AI systems."
    )

    doc.save(output_path)
    print(f"Scalability document successfully created at: {output_path}")

if __name__ == "__main__":
    proj_out = "/Users/shaikmohammedafrazhussain/.gemini/antigravity-ide/scratch/financial-reconciliation-agent/future_scalability_plan.docx"
    artifact_out = "/Users/shaikmohammedafrazhussain/.gemini/antigravity-ide/brain/255a685a-4b23-4b66-b4d9-b2bbde53f572/future_scalability_plan.docx"
    create_scalability_docx(proj_out)
    create_scalability_docx(artifact_out)
